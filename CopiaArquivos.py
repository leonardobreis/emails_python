import os
import re
import zipfile
from collections import defaultdict
from datetime import datetime

pasta_saida = os.path.join(os.path.expanduser("~"), "Downloads")

pastas = {
    "ETIQUETAS": {"pasta": r"V:\DBCorp\Engenharia\Etiquetas", "extensao": "PDF", "nome": "Etiqueta"},
    "PDF": {"pasta": r"V:\DBCorp\Engenharia\PDF", "extensao": "PDF", "nome": "PDF"},
    "DXF": {"pasta": r"V:\DBCorp\Engenharia\DXF", "extensao": "DXF", "nome": "DXF"},
    "STEP": {"pasta": r"V:\DBCorp\Engenharia\STEP", "extensao": "STEP", "nome": "STEP"}
}

ordem_tipos = ["PDF", "ETIQUETAS", "DXF", "STEP"]


def formatar_codigo(codigo):
    return f"{codigo[:2]}.{codigo[2:4]}.{codigo[4:6]}.{codigo[6:]}"


def normalizar_codigo_entrada(valor):
    valor = valor.strip()
    if not re.fullmatch(r"[\d.\-_\s]+", valor):
        return None

    codigo = re.sub(r"\D", "", valor)
    return codigo if len(codigo) == 9 else None


def extrair_codigo_nome(nome):
    stem = os.path.splitext(os.path.basename(nome))[0]

    match = re.match(r"^\s*(\d{9})(?=\D|$)", stem)
    if match:
        return match.group(1)

    match = re.match(
        r"^\s*(\d{2})\s*[.\-_\s]\s*(\d{2})\s*[.\-_\s]\s*(\d{2})\s*[.\-_\s]\s*(\d{3})(?=\D|$)",
        stem
    )

    return "".join(match.groups()) if match else None


def extrair_revisao(nome):
    stem = os.path.splitext(os.path.basename(nome))[0]
    matches = list(re.finditer(r"(?i)\bREV\s*[-._:]?\s*(\d{1,3})(?=\D|$)", stem))

    if len(matches) == 1:
        return int(matches[0].group(1)), None

    if len(matches) > 1:
        return None, "Mais de uma revisão encontrada no nome"

    if re.search(r"(?i)\bREV", stem):
        return None, "Revisão não identificada com segurança"

    return None, "Marcador REV não encontrado"


def sintaxe_toleravel(nome, extensao_esperada):
    codigo = r"(?:\d{9}|\d{2}\s*[.\-_\s]\s*\d{2}\s*[.\-_\s]\s*\d{2}\s*[.\-_\s]\s*\d{3})"
    padrao = rf"^\s*{codigo}\s*(?:-\s*)?REV\s*[-._:]?\s*\d{{1,3}}\s*\.{re.escape(extensao_esperada)}\s*$"

    return bool(re.fullmatch(padrao, nome, re.IGNORECASE))


def analisar_arquivo(tipo, caminho):
    nome = os.path.basename(caminho)
    codigo = extrair_codigo_nome(nome)
    revisao, erro_revisao = extrair_revisao(nome)
    extensao = os.path.splitext(nome)[1][1:]
    extensao_esperada = pastas[tipo]["extensao"]
    problemas = []

    if codigo is None:
        problemas.append("Código não identificado com segurança")

    if erro_revisao:
        problemas.append(erro_revisao)

    extensao_valida = extensao.upper() == extensao_esperada.upper()

    if not extensao_valida:
        problemas.append(
            f"Extensão .{extensao or '(sem extensão)'} diferente da esperada .{extensao_esperada}"
        )

    estrutura_segura = (
        codigo is not None
        and revisao is not None
        and extensao_valida
        and sintaxe_toleravel(nome, extensao_esperada)
    )

    canonico = False

    if codigo is not None and revisao is not None and extensao_valida:
        esperado = f"{formatar_codigo(codigo)} - REV {revisao:02d}.{extensao_esperada}"
        canonico = nome.casefold() == esperado.casefold()

        if not estrutura_segura:
            problemas.append("Estrutura do nome não reconhecida com segurança")
        elif not canonico:
            problemas.append(f"Nomenclatura fora do padrão. Esperado: {esperado}")

    return {
        "tipo": tipo,
        "nome_tipo": pastas[tipo]["nome"],
        "caminho": caminho,
        "nome": nome,
        "codigo": codigo,
        "revisao": revisao,
        "extensao_valida": extensao_valida,
        "estrutura_segura": estrutura_segura,
        "canonico": canonico,
        "problemas": problemas
    }


def carregar_catalogo():
    catalogo = {tipo: defaultdict(list) for tipo in pastas}
    todos_arquivos = []

    for tipo, config in pastas.items():
        pasta = config["pasta"]

        if not os.path.isdir(pasta):
            raise FileNotFoundError(f"Pasta não encontrada ou inacessível: {pasta}")

        with os.scandir(pasta) as entradas:
            for entrada in entradas:
                if not entrada.is_file() or entrada.name.lower() in {"thumbs.db", "desktop.ini"}:
                    continue

                info = analisar_arquivo(tipo, entrada.path)
                todos_arquivos.append(info)

                if info["codigo"]:
                    catalogo[tipo][info["codigo"]].append(info)

    return catalogo, todos_arquivos


def arquivos_com_revisao(infos):
    return [
        info for info in infos
        if info["revisao"] is not None and info["extensao_valida"]
    ]


def maior_revisao(infos):
    validos = arquivos_com_revisao(infos)

    if not validos:
        return None, []

    revisao = max(info["revisao"] for info in validos)

    return revisao, [
        info for info in validos
        if info["revisao"] == revisao
    ]


def maior_revisao_segura(infos):
    validos = [
        info for info in infos
        if info["revisao"] is not None
        and info["extensao_valida"]
        and info["estrutura_segura"]
    ]

    if not validos:
        return None, []

    revisao = max(info["revisao"] for info in validos)

    return revisao, [
        info for info in validos
        if info["revisao"] == revisao
    ]


def gerar_auditoria(catalogo, todos_arquivos):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

    except ImportError:
        print("\n[ERRO] A auditoria em XLSX requer a biblioteca openpyxl.")
        print("Instale com: pip install openpyxl")
        return

    codigos = sorted({
        info["codigo"]
        for info in todos_arquivos
        if info["codigo"]
    })

    resumo = []
    inconsistencias = []

    for codigo in codigos:
        revisoes_tipo = {}
        problemas_resultado = set()

        for tipo in ordem_tipos:
            infos = catalogo[tipo].get(codigo, [])
            revisao, _ = maior_revisao(infos)
            revisoes_tipo[tipo] = revisao

            for info in infos:
                for problema in info["problemas"]:
                    inconsistencias.append([
                        codigo,
                        info["nome_tipo"],
                        info["nome"],
                        problema
                    ])
                    problemas_resultado.add("Nomenclatura")

            grupos = defaultdict(list)

            for info in arquivos_com_revisao(infos):
                grupos[info["revisao"]].append(info)

            for revisao_duplicada, duplicados in grupos.items():
                if len(duplicados) > 1:
                    inconsistencias.append([
                        codigo,
                        pastas[tipo]["nome"],
                        " | ".join(info["nome"] for info in duplicados),
                        f"Duplicidade da REV {revisao_duplicada:02d}"
                    ])
                    problemas_resultado.add("Duplicidade")

        revisoes_existentes = [
            rev for rev in revisoes_tipo.values()
            if rev is not None
        ]

        if len(set(revisoes_existentes)) > 1:
            maior_geral = max(revisoes_existentes)
            problemas_resultado.add("Inconsistência de revisão")

            for tipo in ordem_tipos:
                rev = revisoes_tipo[tipo]

                if rev is not None and rev != maior_geral:
                    inconsistencias.append([
                        codigo,
                        pastas[tipo]["nome"],
                        "",
                        f"REV {rev:02d} diferente da maior revisão encontrada REV {maior_geral:02d}"
                    ])

        if (
            revisoes_tipo["PDF"] is None
            and revisoes_tipo["ETIQUETAS"] is None
            and (
                revisoes_tipo["DXF"] is not None
                or revisoes_tipo["STEP"] is not None
            )
        ):
            problemas_resultado.add("Sem arquivo base")

            inconsistencias.append([
                codigo,
                "Geral",
                "",
                "Possui DXF/STEP, mas não possui PDF nem Etiqueta válidos"
            ])

        resultado = (
            "OK"
            if not problemas_resultado
            else "; ".join(sorted(problemas_resultado))
        )

        resumo.append([
            codigo,
            f"{revisoes_tipo['PDF']:02d}" if revisoes_tipo["PDF"] is not None else "—",
            f"{revisoes_tipo['ETIQUETAS']:02d}" if revisoes_tipo["ETIQUETAS"] is not None else "—",
            f"{revisoes_tipo['DXF']:02d}" if revisoes_tipo["DXF"] is not None else "—",
            f"{revisoes_tipo['STEP']:02d}" if revisoes_tipo["STEP"] is not None else "—",
            resultado
        ])

    for info in todos_arquivos:
        if info["codigo"] is None:
            inconsistencias.append([
                "",
                info["nome_tipo"],
                info["nome"],
                "; ".join(info["problemas"]) or "Nome não reconhecido"
            ])

    wb = Workbook()

    ws = wb.active
    ws.title = "Auditoria"
    ws.append(["Código", "PDF", "Etiqueta", "DXF", "STEP", "Resultado"])

    for linha in resumo:
        ws.append(linha)

    ws2 = wb.create_sheet("Inconsistencias")
    ws2.append(["Código", "Tipo", "Arquivo", "Problema"])

    for linha in inconsistencias:
        ws2.append(linha)

    preenchimento_cabecalho = PatternFill("solid", fgColor="1F4E78")
    fonte_cabecalho = Font(color="FFFFFF", bold=True)
    preenchimento_ok = PatternFill("solid", fgColor="C6EFCE")
    preenchimento_erro = PatternFill("solid", fgColor="FFC7CE")

    for planilha in (ws, ws2):
        for celula in planilha[1]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(horizontal="center")

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        for linha in range(2, planilha.max_row + 1):
            planilha.cell(linha, 1).number_format = "@"

    for linha in range(2, ws.max_row + 1):
        resultado = ws.cell(linha, 6)

        resultado.fill = (
            preenchimento_ok
            if resultado.value == "OK"
            else preenchimento_erro
        )

    for indice, largura in enumerate([14, 10, 12, 10, 10, 40], 1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    for indice, largura in enumerate([14, 14, 70, 70], 1):
        ws2.column_dimensions[get_column_letter(indice)].width = largura

    os.makedirs(pasta_saida, exist_ok=True)

    agora = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = os.path.join(
        pasta_saida,
        f"AuditoriaEngenharia_{agora}.xlsx"
    )

    wb.save(caminho)

    print(f"\n[SUCESSO] Auditoria gerada: {caminho}")
    print(f"Códigos auditados: {len(resumo)}")
    print(f"Inconsistências detalhadas: {len(inconsistencias)}")


def validar_codigo_para_zip(codigo, catalogo):
    encontrados = {}
    inconsistencias = []
    erros_impeditivos = []
    avisos = []

    for tipo in ordem_tipos:
        infos = catalogo[tipo].get(codigo, [])

        inseguros = [
            info for info in infos
            if info["revisao"] is None
            or not info["extensao_valida"]
            or not info["estrutura_segura"]
        ]

        for info in inseguros:
            inconsistencias.append(
                f"{info['nome_tipo']}: {info['nome']} -> "
                + "; ".join(info["problemas"])
            )

        revisao, candidatos = maior_revisao_segura(infos)

        if revisao is None:
            encontrados[tipo] = None
            continue

        if len(candidatos) > 1:
            erros_impeditivos.append(
                f"{pastas[tipo]['nome']}: mais de um arquivo válido encontrado "
                f"para a REV {revisao:02d}: "
                + " | ".join(info["nome"] for info in candidatos)
            )

            encontrados[tipo] = None
            continue

        selecionado = candidatos[0]
        encontrados[tipo] = selecionado

        if not selecionado["canonico"]:
            avisos.append(
                f"{selecionado['nome_tipo']}: '{selecionado['nome']}' "
                f"interpretado como REV {revisao:02d}, mas o nome está fora do padrão"
            )

    if encontrados["PDF"] is None and encontrados["ETIQUETAS"] is None:
        erros_impeditivos.append(
            "Nem Etiqueta nem PDF válidos foram encontrados"
        )

    revisoes = [
        info["revisao"]
        for info in encontrados.values()
        if info is not None
    ]

    if revisoes and len(set(revisoes)) > 1:
        maior_geral = max(revisoes)
        detalhes = []

        for tipo in ordem_tipos:
            info = encontrados[tipo]

            if info is not None:
                status = (
                    ""
                    if info["revisao"] == maior_geral
                    else " (DESATUALIZADO)"
                )

                detalhes.append(
                    f"{info['nome_tipo']}: REV {info['revisao']:02d}{status}"
                )

        inconsistencias.append(
            f"Arquivos com revisões diferentes. "
            f"Maior revisão encontrada: REV {maior_geral:02d}. "
            + " | ".join(detalhes)
        )

    return encontrados, inconsistencias, erros_impeditivos, avisos


def gerar_zip(codigos, catalogo, pedido=""):
    arquivos_para_zipar = []
    resumo = {}
    inconsistencias_gerais = []
    erros_impeditivos_gerais = []
    selecionados_por_codigo = {}

    for codigo in codigos:
        encontrados, inconsistencias, erros_impeditivos, avisos = validar_codigo_para_zip(
            codigo,
            catalogo
        )

        for aviso in avisos:
            print(f"[AVISO] {formatar_codigo(codigo)} - {aviso}")

        if inconsistencias:
            inconsistencias_gerais.append((codigo, inconsistencias))

        if erros_impeditivos:
            erros_impeditivos_gerais.append((codigo, erros_impeditivos))
            continue

        selecionados = [
            info for info in encontrados.values()
            if info is not None
        ]

        if not selecionados:
            erros_impeditivos_gerais.append(
                (codigo, ["Nenhum arquivo válido encontrado"])
            )
            continue

        selecionados_por_codigo[codigo] = selecionados
        revisao = max(info["revisao"] for info in selecionados)

        for info in selecionados:
            arquivos_para_zipar.append(info["caminho"])

        base_resumo = encontrados["PDF"] or encontrados["ETIQUETAS"]

        data_mod = datetime.fromtimestamp(
            os.path.getmtime(base_resumo["caminho"])
        ).strftime("%d/%m/%Y %H:%M")

        resumo[codigo] = (revisao, data_mod)

    if erros_impeditivos_gerais:
        print("\nErros impeditivos:\n")

        for codigo, erros in erros_impeditivos_gerais:
            print(f"[ERRO] {formatar_codigo(codigo)}")

            for erro in erros:
                print(f"       {erro}")

            print()

        print("[AVISO] ZIP não gerado.")
        return

    if inconsistencias_gerais:
        print("\nInconsistências encontradas:\n")

        for codigo, inconsistencias in inconsistencias_gerais:
            print(f"[ATENÇÃO] {formatar_codigo(codigo)}")

            for inconsistencia in inconsistencias:
                print(f"       {inconsistencia}")

            print()

        print("Arquivos que serão incluídos no ZIP:\n")

        for codigo, selecionados in selecionados_por_codigo.items():
            print(formatar_codigo(codigo))

            for info in selecionados:
                print(f"       {info['nome_tipo']}: {info['nome']}")

            print()

        confirmacao = input(
            "Deseja gerar o ZIP mesmo assim? (S/N): "
        ).strip().upper()

        if confirmacao != "S":
            print("\n[AVISO] ZIP não gerado.")
            return

    if not arquivos_para_zipar:
        print("\n[AVISO] Nenhum arquivo válido encontrado para gerar o ZIP.")
        return

    os.makedirs(pasta_saida, exist_ok=True)

    agora = datetime.now().strftime("%Y%m%d_%H%M%S")

    if pedido:
        nome_zip = f"DesenhosAirzap_Ped{pedido}_{agora}.zip"
    else:
        nome_zip = f"DesenhosAirzap_{agora}.zip"

    caminho_zip = os.path.join(pasta_saida, nome_zip)

    try:
        vistos = set()

        with zipfile.ZipFile(
            caminho_zip,
            "w",
            compression=zipfile.ZIP_DEFLATED
        ) as zipf:

            for arquivo in arquivos_para_zipar:
                caminho_normalizado = os.path.normcase(
                    os.path.abspath(arquivo)
                )

                if caminho_normalizado in vistos:
                    continue

                vistos.add(caminho_normalizado)

                zipf.write(
                    arquivo,
                    os.path.basename(arquivo)
                )

        print(f"\n[SUCESSO] ZIP criado: {caminho_zip}")
        print("\nResumo dos arquivos incluídos:")

        for codigo, (revisao, data) in resumo.items():
            print(
                f"  {formatar_codigo(codigo)} "
                f"- REV {revisao:02d} "
                f"- Última alteração: {data}"
            )

    except Exception as e:
        print(f"\n[ERRO] Falha ao criar ZIP: {e}")


def main():
    while True:
        print("\n" + "=" * 60)
        print(f"Destino dos arquivos: {pasta_saida}")
        print("Digite os códigos (Enter duas vezes para finalizar)")
        print('ou digite "auditoria" para auditar todos os arquivos:')
        print('ou digite "sair" para encerrar o programa:')
        print("=" * 60)

        entradas = []
        modo_auditoria = False

        while True:
            entrada = input().strip()

            if entrada.lower() == "sair":
                print("\nPrograma encerrado.")
                return

            if not entrada:
                break

            if not entradas and entrada.lower() == "auditoria":
                modo_auditoria = True
                break

            entradas.append(entrada)

        if modo_auditoria:
            confirmacao = input(
                "Deseja iniciar a auditoria completa dos arquivos de Engenharia? (S/N): "
            ).strip().upper()

            if confirmacao == "S":
                print("\nLendo arquivos de Engenharia...")

                try:
                    catalogo, todos_arquivos = carregar_catalogo()
                    gerar_auditoria(catalogo, todos_arquivos)

                except OSError as e:
                    print(f"\n[ERRO] {e}")

                except Exception as e:
                    print(f"\n[ERRO] Falha durante a auditoria: {e}")

            else:
                print("\n[AVISO] Auditoria cancelada.")

            input("\nPressione Enter para continuar...")
            continue

        if not entradas:
            print("\n[AVISO] Nenhum código informado.")
            input("\nPressione Enter para continuar...")
            continue

        codigos = []
        codigos_invalidos = []

        for entrada in entradas:
            codigo = normalizar_codigo_entrada(entrada)

            if codigo is None:
                codigos_invalidos.append(entrada)

            elif codigo not in codigos:
                codigos.append(codigo)

        if codigos_invalidos:
            print("\n[ERRO] Códigos de entrada inválidos:")

            for valor in codigos_invalidos:
                print(f"       {valor}")

            print("\n[AVISO] ZIP não gerado.")
            input("\nPressione Enter para continuar...")
            continue

        while True:
            pedido = input(
                "\nNúmero do Pedido de Compra (Enter para deixar sem pedido): "
            ).strip()

            if not pedido or pedido.isdigit():
                break

            print(
                "[ERRO] Informe somente o número do pedido ou pressione Enter."
            )

        print("\nLendo arquivos de Engenharia...")

        try:
            catalogo, _ = carregar_catalogo()
            gerar_zip(codigos, catalogo, pedido)

        except OSError as e:
            print(f"\n[ERRO] {e}")

        except Exception as e:
            print(f"\n[ERRO] Falha durante o processamento: {e}")

        #input("\nPressione Enter para continuar...")


if __name__ == "__main__":
    main()